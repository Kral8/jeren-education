#!/usr/bin/env python3
"""Convert source dictionaries to web-optimized searchable JSON."""
from __future__ import annotations

import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Install python-docx: pip install python-docx", file=sys.stderr)
    raise

try:
    import win32com.client
except ImportError:
    win32com = None

SOURCE = Path(r"D:\сайт тещи\словари\Словари")
TARGET = Path(r"D:\сайт тещи\jeren-education\data\dictionaries")
ARCHIVES = TARGET / "archives"
WD_FORMAT_DOCX = 16


def normalize_key(word: str) -> str:
    return word.strip().lower().replace("ё", "е")


def save_entries(entries: dict, filename: str) -> int:
    out = TARGET / filename
    with open(out, "w", encoding="utf-8") as f:
        json.dump(entries, f, ensure_ascii=False, separators=(",", ":"))
    return len(entries)


def word_to_docx(source: Path) -> Path:
    if source.suffix.lower() == ".docx":
        return source
    if win32com is None:
        raise RuntimeError(f"pywin32 required to convert {source.name}")

    tmp = Path(tempfile.gettempdir()) / f"dict-{source.stem}.docx"
    if tmp.exists() and tmp.stat().st_mtime >= source.stat().st_mtime:
        return tmp

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(source.resolve()), ReadOnly=True)
        doc.SaveAs2(str(tmp.resolve()), FileFormat=WD_FORMAT_DOCX)
        doc.Close(False)
    finally:
        word.Quit()
    return tmp


def word_to_text(source: Path, target: Path) -> Path:
    if target.exists() and target.stat().st_mtime >= source.stat().st_mtime:
        return target
    if win32com is None:
        raise RuntimeError(f"pywin32 required to export {source.name}")

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(source.resolve()), ReadOnly=True)
        doc.SaveAs2(str(target.resolve()), FileFormat=7)
        doc.Close(False)
    finally:
        word.Quit()
    return target


def convert_phraseology():
    src = SOURCE / "фразеологический словарь.json"
    entries = {}
    if src.exists():
        with open(src, "r", encoding="utf-8") as f:
            data = json.load(f)
        for slug, item in data.items():
            name = item.get("name", slug)
            key = normalize_key(name)
            entries[key] = {
                "word": name,
                "meaning": item.get("meaning", ""),
                "syntax": item.get("syntax", ""),
                "semanticGroup": item.get("semantic_group", ""),
                "etymology": item.get("etymology", ""),
                "synonyms": item.get("synonyms", []),
                "antonyms": item.get("antonyms", []),
                "examples": item.get("examples", []),
                "components": item.get("components", []),
            }

    archive = ARCHIVES / "frazeologizmy.docx"
    if archive.exists():
        for para in Document(str(archive)).paragraphs:
            text = para.text.strip()
            if not text or "\t" not in text:
                continue
            word, meaning = text.split("\t", 1)
            word, meaning = word.strip(), meaning.strip()
            if not word or not meaning:
                continue
            key = normalize_key(word)
            if key not in entries:
                entries[key] = {"word": word, "meaning": meaning}

    children = ARCHIVES / "veselyy-frazeologicheskiy-dlya-detey.docx"
    if children.exists():
        current = None
        lines: list[str] = []
        for para in Document(str(children)).paragraphs:
            text = para.text.strip()
            if not text:
                if current and lines:
                    key = normalize_key(current)
                    if key not in entries:
                        entries[key] = {"word": current, "meaning": " ".join(lines)}
                current = None
                lines = []
                continue
            if re.match(r"^[А-ЯЁA-Z][А-ЯЁA-Z\s\-]+$", text) and len(text) > 3:
                if current and lines:
                    key = normalize_key(current)
                    if key not in entries:
                        entries[key] = {"word": current, "meaning": " ".join(lines)}
                current = text.title() if text.isupper() else text
                lines = []
            elif current and not text.startswith("-"):
                lines.append(text)
        if current and lines:
            key = normalize_key(current)
            if key not in entries:
                entries[key] = {"word": current, "meaning": " ".join(lines)}

    return save_entries(entries, "phraseology.json")


def convert_foreign_words():
    src = SOURCE / "Словарь иностранных слов .json"
    with open(src, "r", encoding="utf-8") as f:
        data = json.load(f)

    entries = {}
    for word, definition in data.items():
        key = normalize_key(word)
        entries[key] = {"word": word, "meaning": definition}
    return save_entries(entries, "foreign-words.json")


def chunk_key(word: str) -> str:
    w = word.strip().lower()
    if not w:
        return "_other"
    ch = w[0]
    if ch in "абвгдеёжзийклмнопрстуфхцчшщъыьэюя":
        return ch
    if ch == "-":
        return "_prefix"
    return "_other"


def convert_ozhegov():
    src = SOURCE / "ozhegov.txt"
    chunks: dict[str, dict] = {}

    with open(src, "r", encoding="utf-8", errors="replace") as f:
        f.readline()
        for line in f:
            parts = line.rstrip("\n").split("|")
            if len(parts) < 8:
                continue
            vocab = parts[0].strip()
            if not vocab:
                continue
            key = normalize_key(vocab)
            entry = {
                "word": vocab,
                "baseForm": parts[1].strip(),
                "phonetic": parts[2].strip(),
                "grammar": parts[3].strip(),
                "style": parts[4].strip(),
                "meaning": parts[5].strip(),
                "antonym": parts[6].strip(),
                "examples": parts[7].strip(),
            }
            ck = chunk_key(vocab)
            bucket = chunks.setdefault(ck, {})
            bucket.setdefault(key, []).append(entry)

    oz_dir = TARGET / "ozhegov"
    oz_dir.mkdir(parents=True, exist_ok=True)
    letter_index = {}

    for ck, bucket in sorted(chunks.items()):
        fname = f"{ck}.json"
        out = oz_dir / fname
        with open(out, "w", encoding="utf-8") as f:
            json.dump(bucket, f, ensure_ascii=False, separators=(",", ":"))
        letter_index[ck] = {"file": f"ozhegov/{fname}", "entries": len(bucket)}

    meta = {
        "id": "ozhegov",
        "totalKeys": sum(x["entries"] for x in letter_index.values()),
        "chunks": letter_index,
    }
    with open(oz_dir / "meta.json", "w", encoding="utf-8") as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

    return meta["totalKeys"]


def convert_paronims():
    src = ARCHIVES / "paronimy.docx"
    entries: dict[str, dict] = {}
    for para in Document(str(src)).paragraphs:
        text = para.text.strip()
        if not text or re.match(r"^-?\d+-?$", text):
            continue
        if re.search(r"\s[–—-]\s", text) and len(text) < 100 and not text.endswith("."):
            parts = re.split(r"\s[–—-]\s", text, 1)
            word = parts[0].strip()
            pair = parts[1].strip() if len(parts) > 1 else ""
            if word:
                key = normalize_key(word.split(".")[0])
                entries[key] = {"word": word, "meaning": f"Пароним: {pair}" if pair else pair, "pair": pair}
            continue
        match = re.match(r"^([А-ЯЁ][а-яё\-]+)\.\s+(.+)$", text)
        if match:
            word, meaning = match.group(1).strip(), match.group(2).strip()
            key = normalize_key(word)
            entry = entries.get(key, {"word": word, "meaning": meaning})
            entry["meaning"] = meaning
            if "pair" in entries.get(key, {}):
                entry["pair"] = entries[key]["pair"]
            entries[key] = entry
    return save_entries(entries, "paronyms.json")


def convert_synonyms():
    src = ARCHIVES / "sinonimy.docx"
    entries: dict[str, dict] = {}
    skip = {
        "от редактора",
        "о словаре",
        "состав словника словаря",
        "имена существительные",
        "имена прилагательные",
        "наречия",
        "структура словаря и словарной статьи",
    }
    for para in Document(str(src)).paragraphs:
        text = para.text.strip()
        if not text or text.lower() in skip:
            continue
        match = re.match(r"^([А-ЯЁ][А-ЯЁ\s\-]{2,45}),\s*(.+)$", text)
        if not match:
            continue
        word = re.sub(r"\s+", " ", match.group(1).strip())
        meaning = match.group(2).strip()
        if len(word) < 2 or len(meaning) < 2:
            continue
        key = normalize_key(word)
        entries[key] = {"word": word.title() if word.isupper() else word, "meaning": meaning}
    return save_entries(entries, "synonyms.json")


def convert_proverbs():
    src = ARCHIVES / "russkie-poslovitsy-i-pogovorki.docx"
    entries: dict[str, dict] = {}
    skip_fragments = ("пословиц", "состав", "редактор", "§", "соч.", "изд.", "гл.", "стр.")
    for para in Document(str(src)).paragraphs:
        text = para.text.strip()
        if not text or len(text) < 12 or len(text) > 260:
            continue
        lower = text.lower()
        if any(x in lower for x in skip_fragments):
            continue
        if not re.search(r"\s[–—-]\s", text):
            continue
        parts = re.split(r"\s[–—-]\s", text, 1)
        if len(parts[0]) < 8:
            continue
        proverb = text
        gloss = parts[1].strip() if len(parts) > 1 and len(parts[1]) > 6 else ""
        key = normalize_key(parts[0][:80])
        if key not in entries:
            entries[key] = {"word": proverb, "meaning": gloss}
    return save_entries(entries, "proverbs.json")


def convert_medieval():
    src = ARCHIVES / "ustarevshie-slova-srednevekovoy-rusi.docx"
    entries: dict[str, dict] = {}
    text = "\n".join(p.text for p in Document(str(src)).paragraphs)
    for chunk in re.split(r"[\n\r]+", text):
        chunk = chunk.strip()
        if not chunk:
            continue
        for sep in (" — ", " – ", " - "):
            if sep not in chunk:
                continue
            word, meaning = chunk.split(sep, 1)
            word = word.strip().strip(".")
            meaning = meaning.strip().rstrip(".")
            if word and meaning and len(word) <= 60:
                entries[normalize_key(word)] = {"word": word, "meaning": meaning}
            break
        else:
            inline = re.findall(
                r"([А-ЯЁA-Z][а-яёa-z\-]{1,40})\s*[—–-]\s*([^\.]+?)(?=\.[А-ЯЁA-Z]|\.$|$)",
                chunk,
            )
            for word, meaning in inline:
                word, meaning = word.strip(), meaning.strip()
                if word and meaning:
                    entries[normalize_key(word)] = {"word": word, "meaning": meaning}
    return save_entries(entries, "medieval-words.json")


def convert_antonyms():
    src = ARCHIVES / "lvov-antonimy-1984.doc"
    docx = word_to_docx(src)
    entries: dict[str, dict] = {}
    for para in Document(str(docx)).paragraphs:
        text = para.text.strip()
        if not text:
            continue
        pair_match = re.match(r"^([а-яё\-]+)\s*[—–-]\s*([а-яё\-]+)$", text, re.I)
        if pair_match:
            word = pair_match.group(1).strip()
            antonym = pair_match.group(2).strip()
            key = normalize_key(word)
            entries[key] = {"word": word, "meaning": f"Антоним: {antonym}", "antonym": antonym}
            continue
        tab_match = re.match(r"^\([^)]+\)\s*=\s*(.+)$", text)
        if tab_match and "—" in tab_match.group(1):
            entries[normalize_key(tab_match.group(1)[:40])] = {
                "word": tab_match.group(1).strip(),
                "meaning": "Антонимическая пара",
            }
    return save_entries(entries, "antonyms-lvov.json")


def convert_efremova():
    src = ARCHIVES / "efremova-2000.doc"
    txt_path = TARGET / "_efremova-export.txt"
    word_to_text(src, txt_path)
    entries: dict[str, dict] = {}
    headword_re = re.compile(r"^[а-яё][а-яё\-]*$")
    current_word = None
    current_lines: list[str] = []

    def flush():
        nonlocal current_word, current_lines
        if current_word and current_lines:
            key = normalize_key(current_word)
            if key not in entries:
                entries[key] = {
                    "word": current_word,
                    "meaning": " ".join(current_lines).strip(),
                }
        current_word = None
        current_lines = []

    with open(txt_path, "r", encoding="utf-16") as f:
        for line in f:
            stripped = line.strip()
            if not stripped:
                flush()
                continue
            if headword_re.match(stripped) and len(stripped) >= 2 and not stripped.startswith("-"):
                flush()
                current_word = stripped
                current_lines = []
            elif current_word:
                current_lines.append(stripped)

    flush()
    return save_entries(entries, "efremova.json")


def copy_archives():
    ARCHIVES.mkdir(parents=True, exist_ok=True)
    archive_defs = [
        ("Веселый фразеологический словарь для детей.docx", "veselyy-frazeologicheskiy-dlya-detey.docx"),
        ("Львов М.Р. Словарь антонимов русского языка (1984).doc", "lvov-antonimy-1984.doc"),
        ("РУССКИЕ ПОСЛОВИЦЫ И ПОГОВОРКИ.docx", "russkie-poslovitsy-i-pogovorki.docx"),
        ("СЛОВАРЬ УСТАРЕВШИХ СЛОВ СРЕДНЕВЕКОВОЙ РУСИ.docx", "ustarevshie-slova-srednevekovoy-rusi.docx"),
        ("Словарь паронимов   Краткий словарь паронимов русского языка.docx", "paronimy.docx"),
        ("Словарь синонимов русского языка Под общей ред.docx", "sinonimy.docx"),
        ("Словарь фразеологизмов.docx", "frazeologizmy.docx"),
        ("Т.Ф. Ефремова - Новый толково-словообразовательный словарь русского языка.2000.doc", "efremova-2000.doc"),
    ]
    copied = []
    for src_name, dst_name in archive_defs:
        src = SOURCE / src_name
        dst = ARCHIVES / dst_name
        if src.exists():
            shutil.copy2(src, dst)
            copied.append(dst_name)
    return copied


def build_index(counts: dict[str, int]) -> dict:
    dictionaries = [
        {
            "id": "ozhegov",
            "title": "Толковый словарь (Ожегов)",
            "description": "Толкования, грамматика и примеры употребления.",
            "language": "ru",
            "type": "explanatory",
            "entryCount": counts["ozhegov"],
            "data": "ozhegov/meta.json",
            "format": "chunked-json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "efremova",
            "title": "Толковый словарь (Ефремова)",
            "description": "Новый толково-словообразовательный словарь русского языка.",
            "language": "ru",
            "type": "explanatory",
            "entryCount": counts["efremova"],
            "data": "efremova.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "foreign-words",
            "title": "Словарь иностранных слов",
            "description": "Иностранные слова, вошедшие в русский язык.",
            "language": "ru",
            "type": "foreign",
            "entryCount": counts["foreign"],
            "data": "foreign-words.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "phraseology",
            "title": "Фразеологический словарь",
            "description": "Фразеологизмы с толкованиями и примерами.",
            "language": "ru",
            "type": "phraseology",
            "entryCount": counts["phraseology"],
            "data": "phraseology.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "paronyms",
            "title": "Словарь паронимов",
            "description": "Парные слова, близкие по звучанию и различные по смыслу.",
            "language": "ru",
            "type": "paronyms",
            "entryCount": counts["paronyms"],
            "data": "paronyms.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "synonyms",
            "title": "Словарь синонимов",
            "description": "Синонимические ряды русского языка.",
            "language": "ru",
            "type": "synonyms",
            "entryCount": counts["synonyms"],
            "data": "synonyms.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "antonyms-lvov",
            "title": "Словарь антонимов (Львов)",
            "description": "Антонимические пары и контекстные примеры.",
            "language": "ru",
            "type": "antonyms",
            "entryCount": counts["antonyms-lvov"],
            "data": "antonyms-lvov.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "proverbs",
            "title": "Русские пословицы и поговорки",
            "description": "Народные пословицы и поговорки с пояснениями.",
            "language": "ru",
            "type": "proverb",
            "entryCount": counts["proverbs"],
            "data": "proverbs.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
        {
            "id": "medieval-words",
            "title": "Устаревшие слова средневековой Руси",
            "description": "Историческая лексика древнерусского периода.",
            "language": "ru",
            "type": "historical",
            "entryCount": counts["medieval-words"],
            "data": "medieval-words.json",
            "format": "json",
            "searchable": True,
            "status": "available",
        },
    ]

    index = {
        "version": 2,
        "generatedFrom": str(SOURCE),
        "totalDictionaries": len(dictionaries),
        "searchableCount": len(dictionaries),
        "dictionaries": dictionaries,
    }

    with open(TARGET / "index.json", "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=2)

    return index


def main():
    TARGET.mkdir(parents=True, exist_ok=True)
    copy_archives()

    text_pdf = ARCHIVES / "text.pdf"
    if text_pdf.exists():
        text_pdf.unlink()

    counts = {
        "phraseology": convert_phraseology(),
        "foreign": convert_foreign_words(),
        "ozhegov": convert_ozhegov(),
        "paronyms": convert_paronims(),
        "synonyms": convert_synonyms(),
        "proverbs": convert_proverbs(),
        "medieval-words": convert_medieval(),
        "antonyms-lvov": convert_antonyms(),
        "efremova": convert_efremova(),
    }
    index = build_index(counts)
    print(json.dumps({"counts": counts, "total": index["totalDictionaries"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
